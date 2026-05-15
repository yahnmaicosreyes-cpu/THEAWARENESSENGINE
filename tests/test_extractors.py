"""Tests for extractors.py — covers helper functions and the three template types."""

import os
import tempfile

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


# ── Shared workbook helpers ───────────────────────────────────────────────────

def _bold_colored(cell, label):
    cell.value = label
    cell.font = Font(bold=True)
    cell.fill = PatternFill(fill_type="solid", fgColor="4472C4")


def _plain_row(ws, row, label, amount=None, daily_col=2, totals_col=5):
    ws.cell(row, 1).value = label
    if amount is not None:
        ws.cell(row, daily_col).value = amount
        ws.cell(row, totals_col).value = amount


def _save_wb(wb):
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)
    return tmp.name


# ── _match_known_category ─────────────────────────────────────────────────────

class TestMatchKnownCategory:
    def test_exact_match(self):
        from extractors import _match_known_category
        assert _match_known_category("Food") == "Food"
        assert _match_known_category("Shelter") == "Shelter"
        assert _match_known_category("Investments") == "Investments"

    def test_case_insensitive(self):
        from extractors import _match_known_category
        assert _match_known_category("food") == "Food"
        assert _match_known_category("FOOD") == "Food"

    def test_alias_match(self):
        from extractors import _match_known_category
        assert _match_known_category("groceries") == "Groceries"
        assert _match_known_category("healthcare") == "Health Care"
        assert _match_known_category("debt repayment") == "Debt Repayment"

    def test_new_da_categories_recognized(self):
        from extractors import _match_known_category
        assert _match_known_category("Tithe") == "Tithe"
        assert _match_known_category("tithe") == "Tithe"
        assert _match_known_category("7th tradition") == "7th Tradition"
        assert _match_known_category("Haircut") == "Haircut"

    def test_unknown_returns_none(self):
        from extractors import _match_known_category
        assert _match_known_category("Car Insurance") is None
        assert _match_known_category("Netflix") is None

    def test_empty_and_none_input(self):
        from extractors import _match_known_category
        assert _match_known_category(None) is None
        assert _match_known_category("") is None


# ── _is_bold_colored_header ───────────────────────────────────────────────────

class TestIsBoldColoredHeader:
    def test_plain_cell_is_not_header(self):
        from extractors import _is_bold_colored_header
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Rent"
        assert _is_bold_colored_header(ws["A1"]) is False

    def test_bold_without_fill_is_not_header(self):
        from extractors import _is_bold_colored_header
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Rent"
        ws["A1"].font = Font(bold=True)
        assert _is_bold_colored_header(ws["A1"]) is False

    def test_bold_with_white_fill_is_not_header(self):
        from extractors import _is_bold_colored_header
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Rent"
        ws["A1"].font = Font(bold=True)
        # Excel stores white as FFFFFFFF (8-char with alpha prefix)
        ws["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
        assert _is_bold_colored_header(ws["A1"]) is False

    def test_bold_colored_is_header(self):
        from extractors import _is_bold_colored_header
        wb = Workbook()
        ws = wb.active
        _bold_colored(ws["A1"], "Housing")
        assert _is_bold_colored_header(ws["A1"]) is True

    def test_empty_cell_is_not_header(self):
        from extractors import _is_bold_colored_header
        wb = Workbook()
        ws = wb.active
        assert _is_bold_colored_header(ws["A1"]) is False


# ── _DA_GUARANTEED ────────────────────────────────────────────────────────────

class TestDAGuaranteed:
    def test_all_24_da_categories_present(self):
        from extractors import _DA_GUARANTEED
        expected = {
            "Tithe", "7th Tradition", "Shelter", "Internet", "Food",
            "Transportation", "Phone Bill", "Laundry", "Storage", "Moving",
            "Gym", "Clothing", "Personal Care", "Haircut", "Health Care",
            "Inner Child", "Entertainment", "Education", "Vacation",
            "Personal Business", "Gifts", "Investments", "Taxes", "Debt Repayment",
        }
        assert expected == _DA_GUARANTEED

    def test_non_da_known_categories_excluded(self):
        from extractors import _DA_GUARANTEED
        assert "Groceries" not in _DA_GUARANTEED
        assert "Dining Out" not in _DA_GUARANTEED
        assert "Utilities" not in _DA_GUARANTEED
        assert "Subscriptions" not in _DA_GUARANTEED


# ── extract_simple ────────────────────────────────────────────────────────────

class TestExtractSimple:
    def test_basic_income_and_categories(self):
        from extractors import extract_simple
        wb = Workbook()
        ws = wb.active
        ws.append(["Income", 3000])
        ws.append(["Rent", 1200])
        ws.append(["Food", 400])
        monthly_data, cats = extract_simple(wb)
        assert monthly_data["BUDGET"]["income"] == 3000
        assert monthly_data["BUDGET"]["Rent"] == 1200
        assert monthly_data["BUDGET"]["Food"] == 400
        assert cats == ["Rent", "Food"]

    def test_structural_words_skipped(self):
        from extractors import extract_simple
        wb = Workbook()
        ws = wb.active
        ws.append(["Income", 3000])
        ws.append(["Necessities", 1500])
        ws.append(["Rent", 1200])
        ws.append(["Subtotal", 1200])
        ws.append(["Total", 4700])
        monthly_data, cats = extract_simple(wb)
        assert "Necessities" not in monthly_data["BUDGET"]
        assert "Subtotal" not in monthly_data["BUDGET"]
        assert "Total" not in monthly_data["BUDGET"]
        assert "Rent" in monthly_data["BUDGET"]

    def test_negative_amounts_captured(self):
        from extractors import extract_simple
        wb = Workbook()
        ws = wb.active
        ws.append(["Income", 3000])
        ws.append(["Groceries", 400])
        ws.append(["Refund", -50])
        monthly_data, _ = extract_simple(wb)
        assert monthly_data["BUDGET"]["Refund"] == -50.0


# ── Daily tracking extraction ─────────────────────────────────────────────────

class TestDailyTrackingExtraction:
    """Verifies the flipped logic: plain sub-rows captured, bold+colored headers skipped."""

    def _make_wb(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "JAN"

        # Row 1: col C = date string (triggers detection), col 5 = TOTALS
        ws.cell(1, 3).value = "2025-01-01"
        ws.cell(1, 5).value = "TOTALS"

        # Income section (plain header + plain sub-row)
        ws.cell(2, 1).value = "INCOME"
        _plain_row(ws, 3, "Salary", 3000)

        # Expense sections: bold+colored section headers, plain sub-rows
        _bold_colored(ws.cell(4, 1), "Housing")
        _plain_row(ws, 5, "Rent", 1200)
        _plain_row(ws, 6, "Electricity", 150)

        _bold_colored(ws.cell(7, 1), "Food & Dining")
        _plain_row(ws, 8, "Groceries", 400)
        _plain_row(ws, 9, "Tithe", 300)    # DA guaranteed
        _plain_row(ws, 10, "Netflix", 15)  # unknown plain category

        return wb

    def test_plain_sub_rows_captured(self):
        from extractors import extract_data
        path = _save_wb(self._make_wb())
        try:
            monthly_data, cats = extract_data(path)
            jan = monthly_data["JAN"]
            assert jan["Rent"] == 1200
            assert jan["Electricity"] == 150
            assert jan["Groceries"] == 400
            assert jan["Tithe"] == 300
            assert jan["Netflix"] == 15
        finally:
            os.unlink(path)

    def test_bold_colored_section_headers_skipped(self):
        from extractors import extract_data
        path = _save_wb(self._make_wb())
        try:
            monthly_data, _ = extract_data(path)
            jan = monthly_data["JAN"]
            assert "Housing" not in jan
            assert "Food & Dining" not in jan
        finally:
            os.unlink(path)

    def test_income_unchanged(self):
        from extractors import extract_data
        path = _save_wb(self._make_wb())
        try:
            monthly_data, _ = extract_data(path)
            assert monthly_data["JAN"]["income"] == 3000
        finally:
            os.unlink(path)

    def test_da_category_captured_when_plain(self):
        from extractors import extract_data
        path = _save_wb(self._make_wb())
        try:
            monthly_data, _ = extract_data(path)
            assert monthly_data["JAN"]["Tithe"] == 300
        finally:
            os.unlink(path)

    def test_unknown_plain_category_captured(self):
        from extractors import extract_data
        path = _save_wb(self._make_wb())
        try:
            monthly_data, _ = extract_data(path)
            assert monthly_data["JAN"]["Netflix"] == 15
        finally:
            os.unlink(path)

    def test_ordered_cats_matches_row_order(self):
        from extractors import extract_data
        path = _save_wb(self._make_wb())
        try:
            _, cats = extract_data(path)
            assert cats.index("Rent") < cats.index("Electricity")
            assert cats.index("Electricity") < cats.index("Groceries")
        finally:
            os.unlink(path)


# ── YYYY Month extraction ─────────────────────────────────────────────────────

class TestYYYYMonthExtraction:
    """Verifies the flipped logic for the YYYY Month template."""

    def _make_wb(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "2025 January"

        # Row 11: detection signature
        ws.cell(11, 1).value = "SPENDING"
        ws.cell(11, 3).value = "Total"

        # Income rows (3-10)
        ws.cell(3, 1).value = "Salary"
        ws.cell(3, 3).value = 3000

        # Expense rows (12+)
        _bold_colored(ws.cell(12, 1), "Housing")
        ws.cell(12, 3).value = 1350   # subtotal — must be skipped

        ws.cell(13, 1).value = "Rent"
        ws.cell(13, 3).value = 1200

        ws.cell(14, 1).value = "Electricity"
        ws.cell(14, 3).value = 150

        _bold_colored(ws.cell(15, 1), "Food")
        ws.cell(15, 3).value = 400    # subtotal — must be skipped

        ws.cell(16, 1).value = "Groceries"
        ws.cell(16, 3).value = 400

        return wb

    def test_plain_rows_captured(self):
        from extractors import extract_data
        path = _save_wb(self._make_wb())
        try:
            monthly_data, _ = extract_data(path)
            jan = monthly_data["JAN"]
            assert jan["Rent"] == 1200
            assert jan["Electricity"] == 150
            assert jan["Groceries"] == 400
        finally:
            os.unlink(path)

    def test_bold_colored_headers_skipped(self):
        from extractors import extract_data
        path = _save_wb(self._make_wb())
        try:
            monthly_data, _ = extract_data(path)
            jan = monthly_data["JAN"]
            assert "Housing" not in jan
            assert "Food" not in jan
        finally:
            os.unlink(path)

    def test_income_unchanged(self):
        from extractors import extract_data
        path = _save_wb(self._make_wb())
        try:
            monthly_data, _ = extract_data(path)
            assert monthly_data["JAN"]["income"] == 3000
        finally:
            os.unlink(path)
